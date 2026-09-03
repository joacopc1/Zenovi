#!/usr/bin/env python3
"""Generate Google Docs-friendly DOCX files from Zenovi's Markdown sources."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "output" / "google-docs"
INK = "000000"
MUTED = "555555"
BORDER = "DADCE0"
LIGHT = "F7F7F7"


def set_font(run, name: str = "Arial", size: float | None = None, bold: bool | None = None,
             italic: bool | None = None, color: str = INK) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths: list[int]) -> None:
    total = sum(widths)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "0")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            set_cell_width(cell, widths[index])
            set_cell_margins(cell)


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), "4")
        tag.set(qn("w:color"), BORDER)


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.first_child_found_in("w:shd")
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def style_paragraph_box(paragraph, *, fill: str | None = None, left_border: str | None = None,
                        bottom_border: str | None = None) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    if fill:
        shading = p_pr.find(qn("w:shd"))
        if shading is None:
            shading = OxmlElement("w:shd")
            p_pr.append(shading)
        shading.set(qn("w:fill"), fill)
    if left_border or bottom_border:
        borders = p_pr.find(qn("w:pBdr"))
        if borders is None:
            borders = OxmlElement("w:pBdr")
            p_pr.append(borders)
        for edge, color, size, space in (
            ("left", left_border, "16", "8"),
            ("bottom", bottom_border, "6", "5"),
        ):
            if not color:
                continue
            node = borders.find(qn(f"w:{edge}"))
            if node is None:
                node = OxmlElement(f"w:{edge}")
                borders.append(node)
            node.set(qn("w:val"), "single")
            node.set(qn("w:sz"), size)
            node.set(qn("w:space"), space)
            node.set(qn("w:color"), color)


def mark_header_row(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = tr_pr.find(qn("w:tblHeader"))
    if header is None:
        header = OxmlElement("w:tblHeader")
        tr_pr.append(header)
    header.set(qn("w:val"), "true")


def add_hyperlink(paragraph, text: str, url: str) -> None:
    part = paragraph.part
    rel_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1155CC")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.extend([color, underline])
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.extend([r_pr, text_node])
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


INLINE_RE = re.compile(r"(\[[^\]]+\]\(https?://[^)]+\)|https?://\S+|\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*)")


def add_inline(paragraph, text: str) -> None:
    cursor = 0
    for match in INLINE_RE.finditer(text):
        if match.start() > cursor:
            set_font(paragraph.add_run(text[cursor:match.start()]), size=11)
        token = match.group(0)
        link_match = re.fullmatch(r"\[([^\]]+)\]\((https?://[^)]+)\)", token)
        if link_match:
            add_hyperlink(paragraph, link_match.group(1), link_match.group(2))
        elif token.startswith("http"):
            add_hyperlink(paragraph, token.rstrip(".,;"), token.rstrip(".,;"))
        elif token.startswith("**"):
            set_font(paragraph.add_run(token[2:-2]), size=11, bold=True)
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_font(run, name="Courier New", size=10, color="333333")
            run.font.highlight_color = None
        elif token.startswith("*"):
            set_font(paragraph.add_run(token[1:-1]), size=11, italic=True)
        cursor = match.end()
    if cursor < len(text):
        set_font(paragraph.add_run(text[cursor:]), size=11)


def set_paragraph_spacing(paragraph, before=0, after=8, line=1.15) -> None:
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.15

    tokens = {
        "Heading 1": (20, INK, 20, 8),
        "Heading 2": (16, INK, 16, 5),
        "Heading 3": (14, "434343", 12, 4),
    }
    for name, (size, color, before, after) in tokens.items():
        style = doc.styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        style = doc.styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.5)
        style.paragraph_format.first_line_indent = Inches(-0.25)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.15


def add_title_block(doc: Document, title: str, subtitle: str) -> None:
    paragraph = doc.add_paragraph()
    set_paragraph_spacing(paragraph, before=0, after=3, line=1.0)
    run = paragraph.add_run(title)
    set_font(run, size=26, bold=True)
    paragraph = doc.add_paragraph()
    set_paragraph_spacing(paragraph, after=10, line=1.15)
    run = paragraph.add_run(subtitle)
    set_font(run, size=11, color=MUTED)
    paragraph = doc.add_paragraph()
    set_paragraph_spacing(paragraph, after=18, line=1.0)
    run = paragraph.add_run("Versión 0.1 · Documento vivo")
    set_font(run, size=9, color=MUTED)


def parse_table(lines: list[str]) -> list[list[str]]:
    rows = []
    for line in lines:
        cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
        if all(re.fullmatch(r":?-{3,}:?", cell or "---") for cell in cells):
            continue
        rows.append(cells)
    return rows


def calculate_widths(rows: list[list[str]], total=9360) -> list[int]:
    count = max(len(row) for row in rows)
    scores = []
    for col in range(count):
        values = [len(row[col]) if col < len(row) else 0 for row in rows]
        scores.append(max(8, min(45, max(values, default=8))))
    raw = [int(total * score / sum(scores)) for score in scores]
    minimum = 900 if count <= 4 else 650
    raw = [max(minimum, value) for value in raw]
    scale = total / sum(raw)
    widths = [int(value * scale) for value in raw]
    widths[-1] += total - sum(widths)
    return widths


def add_table(doc: Document, rows: list[list[str]]) -> None:
    columns = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=columns)
    widths = calculate_widths(rows)
    set_table_geometry(table, widths)
    set_table_borders(table)
    mark_header_row(table.rows[0])
    for row_index, values in enumerate(rows):
        for col_index in range(columns):
            cell = table.cell(row_index, col_index)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            set_paragraph_spacing(paragraph, after=0, line=1.08)
            value = values[col_index] if col_index < len(values) else ""
            add_inline(paragraph, value)
            for run in paragraph.runs:
                set_font(run, size=9.5, bold=(row_index == 0), color="FFFFFF" if row_index == 0 else INK)
            if row_index == 0:
                shade_cell(cell, "202124")
            elif row_index % 2 == 0:
                shade_cell(cell, "F8F9FA")
    after = doc.add_paragraph()
    set_paragraph_spacing(after, after=2, line=1.0)


def add_markdown(doc: Document, source: Path) -> None:
    lines = source.read_text(encoding="utf-8").splitlines()
    paragraph_lines: list[str] = []

    def flush_paragraph() -> None:
        nonlocal paragraph_lines
        if paragraph_lines:
            text = " ".join(item.strip() for item in paragraph_lines).strip()
            if text:
                paragraph = doc.add_paragraph()
                set_paragraph_spacing(paragraph)
                add_inline(paragraph, text)
                if text.startswith("Conectar -> entender -> decidir"):
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    set_paragraph_spacing(paragraph, before=6, after=16, line=1.0)
                    style_paragraph_box(paragraph, fill="F7F7F7", bottom_border=BORDER)
                    for run in paragraph.runs:
                        set_font(run, size=11, bold=True)
            paragraph_lines = []

    index = 0
    while index < len(lines):
        raw = lines[index].rstrip()
        stripped = raw.strip()
        if index == 0 and stripped.startswith("# "):
            index += 1
            continue
        if index < 8 and re.match(r"^(Versión|Fecha|Propietarios|Dominio|Estado|Ventana objetivo|Audiencia):", stripped):
            index += 1
            continue
        if not stripped or stripped == "---":
            flush_paragraph()
            index += 1
            continue
        if stripped.startswith("|") and index + 1 < len(lines) and lines[index + 1].strip().startswith("|"):
            flush_paragraph()
            table_lines = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index])
                index += 1
            add_table(doc, parse_table(table_lines))
            continue
        heading = re.match(r"^(#{2,4})\s+(.+)$", stripped)
        if heading:
            flush_paragraph()
            level = len(heading.group(1)) - 1
            paragraph = doc.add_paragraph(style=f"Heading {level}")
            add_inline(paragraph, heading.group(2))
            if level == 1:
                style_paragraph_box(paragraph, bottom_border=BORDER)
            index += 1
            continue
        if stripped.startswith("> "):
            flush_paragraph()
            paragraph = doc.add_paragraph()
            set_paragraph_spacing(paragraph, before=4, after=12, line=1.15)
            paragraph.paragraph_format.left_indent = Inches(0.25)
            paragraph.paragraph_format.right_indent = Inches(0.12)
            style_paragraph_box(paragraph, fill="F7F7F7", left_border="202124")
            run = paragraph.add_run(stripped[2:])
            set_font(run, size=11, italic=True, color="434343")
            index += 1
            continue
        bullet = re.match(r"^(\s*)[-*]\s+(.+)$", raw)
        numbered = re.match(r"^(\s*)\d+\.\s+(.+)$", raw)
        if bullet or numbered:
            flush_paragraph()
            match = bullet or numbered
            assert match is not None
            level = min(2, len(match.group(1).replace("\t", "  ")) // 2)
            paragraph = doc.add_paragraph(style="List Bullet" if bullet else "List Number")
            paragraph.paragraph_format.left_indent = Inches(0.5 + level * 0.25)
            paragraph.paragraph_format.first_line_indent = Inches(-0.25)
            text = match.group(2).replace("[ ]", "☐").replace("[x]", "☒")
            add_inline(paragraph, text)
            index += 1
            continue
        paragraph_lines.append(stripped)
        index += 1
    flush_paragraph()


def build(source: Path, destination: Path, title: str, subtitle: str) -> None:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    configure_styles(doc)
    add_title_block(doc, title, subtitle)
    add_markdown(doc, source)
    properties = doc.core_properties
    properties.title = title
    properties.subject = subtitle
    properties.author = "Zenovi"
    properties.keywords = "Zenovi, SaaS, producto, MVP"
    destination.parent.mkdir(parents=True, exist_ok=True)
    doc.save(destination)


def main() -> None:
    OUTPUT.mkdir(parents=True, exist_ok=True)
    build(
        ROOT / "docs" / "PRD.md",
        OUTPUT / "Zenovi-PRD-v0.1.docx",
        "Zenovi — Product Requirements Document",
        "Visión completa del producto, requisitos, arquitectura y criterios de lanzamiento.",
    )
    build(
        ROOT / "docs" / "MVP-SPEC.md",
        OUTPUT / "Zenovi-MVP-v0.1.docx",
        "Zenovi — MVP",
        "Alcance, criterios de aceptación y plan de ejecución para la primera versión.",
    )
    build(
        ROOT / "docs" / "research" / "icp-avatar-research.md",
        OUTPUT / "Zenovi-ICP-Avatar-v0.1.docx",
        "Zenovi — Investigación de ICP y avatar",
        "Hipótesis, problemas, criterios de selección y plan de validación del mercado inicial.",
    )
    print(f"Generated Google Docs-compatible files in {OUTPUT}")


if __name__ == "__main__":
    main()
